"""
Генерация docx из шаблонов.
Приоритет: шаблон из БД (DocTemplate) → файл из templates_docx/
"""
import os, re
from io import BytesIO
from docx import Document

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), 'templates_docx')


def _get_template_path(template_type):
    """Возвращает путь к шаблону: сначала ищет в БД, потом в папке."""
    try:
        from .models import DocTemplate
        tpl = DocTemplate.objects.filter(template_type=template_type).first()
        if tpl and tpl.file:
            return tpl.file.path
    except Exception:
        pass
    return os.path.join(TEMPLATES_DIR, f'{template_type}.docx')


def _replace_in_paragraph(para, replacements):
    full_text = para.text
    if '{{' not in full_text:
        return
    new_text = full_text
    for key, value in replacements.items():
        new_text = new_text.replace('{{' + key + '}}', str(value) if value is not None else '')
    if new_text == full_text:
        return
    if para.runs:
        para.runs[0].text = new_text
        for run in para.runs[1:]:
            run.text = ''
    else:
        para.add_run(new_text)


def _replace_in_doc(doc, replacements):
    for para in doc.paragraphs:
        _replace_in_paragraph(para, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para, replacements)


def _month_name(month_num):
    months = ['','января','февраля','марта','апреля','мая','июня',
              'июля','августа','сентября','октября','ноября','декабря']
    try:
        return months[int(month_num)]
    except (ValueError, IndexError):
        return str(month_num)


def generate_harakteristika(char_doc):
    path = _get_template_path('harakteristika')
    doc  = Document(path)

    replacements = {
        'fio':               char_doc.fio or '',
        'library':           char_doc.library or '',
        'number':            char_doc.work_days_count or '',
        'not_day_one':       char_doc.missed_days or '0',
        'not_day_two':       char_doc.missed_unexcused or '0',
        'special':           char_doc.speciality or '',
        'good':              char_doc.work_quality or '',
        'boss_organization': char_doc.supervisor or '',
    }

    if char_doc.work_start:
        replacements.update({
            'day_begin':   char_doc.work_start.day,
            'month_begin': _month_name(char_doc.work_start.month),
            'year_begin':  char_doc.work_start.year,
        })
    else:
        replacements.update({'day_begin': '', 'month_begin': '', 'year_begin': ''})

    if char_doc.work_end:
        replacements.update({
            'day_finish':   char_doc.work_end.day,
            'month_finish': _month_name(char_doc.work_end.month),
            'year_finish':  char_doc.work_end.year,
        })
    else:
        replacements.update({'day_finish': '', 'month_finish': '', 'year_finish': ''})

    _replace_in_doc(doc, replacements)
    out = BytesIO(); doc.save(out); out.seek(0)
    return out


def _practice_type_ru(practice_type):
    return {'educational': 'учебную', 'production': 'производственную',
            'prediploma': 'преддипломную'}.get(practice_type, practice_type)


def generate_attestat(practice_doc):
    tpl_key = ('attestat_proizvodstvennaya'
               if practice_doc.practice_type in ('production', 'prediploma')
               else 'attestat_uchebnaya')

    path = _get_template_path(tpl_key)
    doc  = Document(path)

    sd = practice_doc.start_date
    ed = practice_doc.end_date

    replacements = {
        'fio':   practice_doc.fio or '',
        'spec':  practice_doc.speciality or '',
        'grupa': practice_doc.group or '',
        'kurs':  practice_doc.course or '',
        'obuch': practice_doc.education_form or '',
        'data':  sd.day   if sd else '',
        'data2': _month_name(sd.month) if sd else '',
        'god':   sd.year  if sd else '',
        'data3': ed.day   if ed else '',
        'data4': _month_name(ed.month) if ed else '',
        'god1':  ed.year  if ed else '',
        'vid':   _practice_type_ru(practice_doc.practice_type),
        'kod':   practice_doc.module_code or '',
        'mesto': practice_doc.organization or '',
        'adress':practice_doc.organization_address or '',
        'ruka':  practice_doc.supervisor or '',
        'da':    '+',
    }
    _replace_in_doc(doc, replacements)

    # Заполняем таблицы плюсами
    quality_col   = {'high': 2, 'medium': 3, 'low': 4}
    competence_col = {'full': 1, 'partial': 2, 'none': 3}
    manifest_col  = {'regular': 2, 'episodic': 3, 'none': 4}

    work_q  = [getattr(practice_doc, f'work_{i}_quality', '') for i in range(1, 12)]
    pk_lvl  = [getattr(practice_doc, f'pk_{i}_level',    '') for i in range(1, 7)]
    ok_lvl  = [getattr(practice_doc, f'ok_{i}_level',    '') for i in range(1, 12)]

    tables = doc.tables

    def _set_plus(table, row_idx, col_idx):
        try:
            cell = table.rows[row_idx].cells[col_idx]
            if not cell.text.strip():
                cell.paragraphs[0].add_run('+')
        except IndexError:
            pass

    if len(tables) > 0:
        for i, q in enumerate(work_q):
            if q in quality_col:
                _set_plus(tables[0], i + 2, quality_col[q])

    if len(tables) > 1:
        for i, lvl in enumerate(pk_lvl):
            if lvl in competence_col:
                _set_plus(tables[1], i + 2, competence_col[lvl])

    if len(tables) > 2:
        for i, lvl in enumerate(ok_lvl):
            if lvl in manifest_col:
                _set_plus(tables[2], i + 2, manifest_col[lvl])

    out = BytesIO(); doc.save(out); out.seek(0)
    return out
