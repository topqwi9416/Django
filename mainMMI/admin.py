from django.contrib import admin
from django.contrib.auth.models import User
from django.contrib.auth.admin import UserAdmin as BaseUserAdmin
from django.utils.html import format_html
from django.urls import reverse
from django.db.models import Count
from .models import UserProfile, PracticeDocument, CharacteristicDocument, DocTemplate


# ─── Inline профиля внутри User ──────────────────────────────────────────────

class UserProfileInline(admin.StackedInline):
    model = UserProfile
    can_delete = False
    verbose_name = 'Профиль'
    verbose_name_plural = 'Профиль'
    fields = (
        ('last_name', 'first_name', 'middle_name'),
        ('phone', 'telegram'),
        ('birth_date', 'gender'),
        'inn',
    )


# ─── Расширенный список пользователей ────────────────────────────────────────

class UserAdmin(BaseUserAdmin):
    inlines = (UserProfileInline,)

    list_display  = ('username', 'get_fio', 'email', 'get_phone',
                     'get_docs_count', 'is_active', 'date_joined')
    list_filter   = ('is_active', 'is_staff', 'date_joined')
    search_fields = ('username', 'email',
                     'profile__last_name', 'profile__first_name', 'profile__phone')
    ordering      = ('-date_joined',)

    @admin.display(description='ФИО')
    def get_fio(self, obj):
        try:
            p = obj.profile
            parts = [p.last_name, p.first_name, p.middle_name]
            fio = ' '.join(x for x in parts if x)
            return fio or '—'
        except UserProfile.DoesNotExist:
            return '—'

    @admin.display(description='Телефон')
    def get_phone(self, obj):
        try:
            return obj.profile.phone or '—'
        except UserProfile.DoesNotExist:
            return '—'

    @admin.display(description='Документов')
    def get_docs_count(self, obj):
        try:
            pd = obj.practice_documents.count()
            ch = obj.characteristics.count()
            total = pd + ch
            return f'{total} (атт: {pd}, хар: {ch})'
        except Exception:
            return '0'


admin.site.unregister(User)
admin.site.register(User, UserAdmin)


# ─── Аттестационные листы ────────────────────────────────────────────────────

@admin.register(PracticeDocument)
class PracticeDocumentAdmin(admin.ModelAdmin):
    list_display  = ('fio', 'get_practice_type', 'group', 'course',
                     'get_dates', 'user', 'created_at')
    list_filter   = ('practice_type', 'created_at')
    search_fields = ('fio', 'speciality', 'group', 'user__username')
    ordering      = ('-created_at',)
    readonly_fields = ('created_at', 'updated_at')

    fieldsets = (
        ('Студент', {
            'fields': ('user', 'fio', 'speciality', 'group', 'course', 'education_form')
        }),
        ('Практика', {
            'fields': ('practice_type', 'start_date', 'end_date',
                       'module_code', 'organization', 'organization_address', 'supervisor')
        }),
        ('Результат', {
            'fields': ('result', 'document_date')
        }),
        ('Служебное', {
            'classes': ('collapse',),
            'fields': ('created_at', 'updated_at')
        }),
    )

    @admin.display(description='Вид практики')
    def get_practice_type(self, obj):
        colors = {
            'educational': '#2f81f7',
            'production':  '#d29922',
            'prediploma':  '#8957e5',
        }
        color = colors.get(obj.practice_type, '#8b949e')
        return format_html(
            '<span style="color:{};font-weight:600">{}</span>',
            color, obj.get_practice_type_display()
        )

    @admin.display(description='Период')
    def get_dates(self, obj):
        if obj.start_date and obj.end_date:
            return f'{obj.start_date.strftime("%d.%m.%Y")} — {obj.end_date.strftime("%d.%m.%Y")}'
        return '—'


# ─── Характеристики ──────────────────────────────────────────────────────────

@admin.register(CharacteristicDocument)
class CharacteristicDocumentAdmin(admin.ModelAdmin):
    list_display  = ('fio', 'speciality', 'get_dates', 'work_days_count', 'user', 'created_at')
    list_filter   = ('created_at',)
    search_fields = ('fio', 'speciality', 'user__username')
    ordering      = ('-created_at',)
    readonly_fields = ('created_at', 'updated_at')

    @admin.display(description='Период работы')
    def get_dates(self, obj):
        if obj.work_start and obj.work_end:
            return f'{obj.work_start.strftime("%d.%m.%Y")} — {obj.work_end.strftime("%d.%m.%Y")}'
        return '—'


# ─── Шаблоны документов ──────────────────────────────────────────────────────

@admin.register(DocTemplate)
class DocTemplateAdmin(admin.ModelAdmin):
    list_display  = ('get_type_badge', 'get_file_link', 'uploaded_by', 'uploaded_at')
    list_filter   = ('template_type',)
    readonly_fields = ('uploaded_at', 'uploaded_by', 'get_file_preview')

    fieldsets = (
        ('Шаблон', {
            'description': (
                'Загрузите .docx файл с плейсхолдерами. '
                'Используйте формат {{имя_переменной}}. '
                'Например: {{fio}}, {{data}}, {{spec}}'
            ),
            'fields': ('template_type', 'file', 'description')
        }),
        ('Информация', {
            'classes': ('collapse',),
            'fields': ('uploaded_by', 'uploaded_at', 'get_file_preview')
        }),
    )

    def save_model(self, request, obj, form, change):
        obj.uploaded_by = request.user
        super().save_model(request, obj, form, change)

    @admin.display(description='Тип шаблона')
    def get_type_badge(self, obj):
        colors = {
            'harakteristika':             ('#3fb950', '📋 Характеристика'),
            'attestat_uchebnaya':         ('#2f81f7', '📝 Атт. лист УП'),
            'attestat_proizvodstvennaya': ('#d29922', '🏭 Атт. лист ПП'),
        }
        color, label = colors.get(obj.template_type, ('#8b949e', obj.get_template_type_display()))
        return format_html(
            '<span style="color:{};font-weight:600">{}</span>', color, label
        )

    @admin.display(description='Файл')
    def get_file_link(self, obj):
        if obj.file:
            return format_html(
                '<a href="{}" style="color:#2f81f7" download>⬇ {}</a>',
                obj.file.url, obj.file.name.split('/')[-1]
            )
        return format_html('<span style="color:#8b949e">— файл не загружен —</span>')

    @admin.display(description='Предпросмотр')
    def get_file_preview(self, obj):
        if not obj.file:
            return '—'
        try:
            from docx import Document as DocxDoc
            import re
            doc = DocxDoc(obj.file.path)
            placeholders = set()
            for para in doc.paragraphs:
                placeholders.update(re.findall(r'\{\{(\w+)\}\}', para.text))
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            placeholders.update(re.findall(r'\{\{(\w+)\}\}', para.text))
            if placeholders:
                chips = ' '.join(
                    f'<code style="background:#161b22;border:1px solid #21262d;'
                    f'border-radius:4px;padding:2px 6px;font-size:12px;color:#e6edf3">'
                    f'{{{{{p}}}}}</code>'
                    for p in sorted(placeholders)
                )
                return format_html(
                    '<div style="margin-top:8px"><strong>Плейсхолдеры в файле:</strong><br>{}</div>',
                    chips
                )
            return 'Плейсхолдеры не найдены'
        except Exception as e:
            return f'Ошибка чтения файла: {e}'


# ─── Настройка сайта админки ─────────────────────────────────────────────────

admin.site.site_header  = 'ПрактикаДок — Администрирование'
admin.site.site_title   = 'ПрактикаДок'
admin.site.index_title  = 'Управление системой'
